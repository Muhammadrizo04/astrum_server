import time
from django.db import models
import qrcode
from io import BytesIO
from django.core.files import File
from PIL import Image, ImageDraw, ImageFont
from django.conf import settings
import os
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from pptx import Presentation
from django.core.files.base import ContentFile


def get_last_sertificate_id():
    last_ids = []
    for model in [PHPModule, PythonModule, ReactModule, CSharpModule]:
        last_obj = model.objects.order_by('-sertificate_id').first()
        if last_obj and last_obj.sertificate_id:
            last_ids.append(int(last_obj.sertificate_id))
    return max(last_ids) if last_ids else 0

def get_last_sertificate_id_numeric():
    last_ids = []
    for model in [PHPModule, PythonModule, ReactModule, CSharpModule]:
        last_obj = model.objects.order_by('-sertificate_id_numeric').first()
        if last_obj and last_obj.sertificate_id_numeric:
            last_ids.append(last_obj.sertificate_id_numeric)
    return max(last_ids) if last_ids else 0

class PHPModule(models.Model):
    first_name = models.CharField(max_length=50)
    last_name = models.CharField(max_length=50)
    middle_name = models.CharField(max_length=50)
    issue_date = models.CharField(max_length=30, blank=True)
    created_at = models.DateTimeField(auto_now_add=True)
    module_1 = models.CharField(max_length=5)
    module_2 = models.CharField(max_length=5)
    module_3 = models.CharField(max_length=5)
    module_4 = models.CharField(max_length=5)
    module_5 = models.CharField(max_length=5)
    certificate_link = models.CharField(max_length=250, blank=True)
    seria = models.CharField(max_length=3, default="PC")
    course = models.CharField(max_length=3, default="PHP")
    sertificate_id = models.CharField(max_length=7, unique=True, null=True, blank=True)
    sertificate_id_numeric = models.IntegerField(unique=True, null=True, blank=True)
    qr_code = models.ImageField(upload_to="module_qr_code", null=True, blank=True)
    pptx_file = models.FileField(upload_to="pptx_module", blank=True)
    sertificate_front = models.FileField(upload_to="module_front", blank=True)
    sertificate_back = models.FileField(upload_to="module_back", blank=True)

    def __str__(self):
        return self.first_name

    @staticmethod
    def add_image(prs, slide_index, image_path, left, top, height):
        slide = prs.slides[slide_index]
        image = Image.open(image_path)
        image = image.convert('RGBA')
        data = list(image.getdata())
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        new_image = Image.new('RGBA', image.size)
        new_image.putdata(new_data)
        image_stream = BytesIO()
        new_image.save(image_stream, format="PNG")
        image_stream.seek(0)
        slide.shapes.add_picture(image_stream, left, top, height=height)

    @staticmethod
    def add_text(prs, slide_index, left, top, width, height, input_text, font_size, font_color,
                 font_name='Gilroy', alignment=PP_ALIGN.LEFT):
        slide = prs.slides[slide_index]
        text_box = slide.shapes.add_textbox(left, top, width, height)
        text_frame = text_box.text_frame
        text = text_frame.add_paragraph()
        text.text = input_text
        text.font.size = Pt(font_size)
        text.font.name = font_name
        text.font.color.rgb = RGBColor(font_color[0], font_color[1], font_color[2])
        text.alignment = alignment
        for run in text.runs:
            run.font.bold = True

    def generate_certificate(self):
        pptx_template_path = os.path.join(settings.MEDIA_ROOT, 'template/php.pptx')
        prs = Presentation(pptx_template_path)

        # Slide 0

        black_color = (0, 0, 0,)
        text_color = (0x54, 0x30, 0xCE)
        text = f"{self.last_name} {self.first_name} {self.middle_name}"
        seria = f"{self.seria} {self.sertificate_id}"
        qr_code = os.path.join(settings.MEDIA_ROOT, f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
        qr_code_img_back = os.path.join(settings.MEDIA_ROOT, f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
        issued = f"{self.issue_date}"

        self.add_image(prs, 0, qr_code, Inches(0.3976377953), Inches(4.7952755906), Inches(1.3))
        self.add_text(prs, 0, Inches(1), Inches(2.55), Inches(8), Inches(1), text, 28, text_color,
                      alignment=PP_ALIGN.CENTER)
        self.add_text(prs, 0, Inches(4.645669), Inches(5.64), Inches(1), Inches(1), issued, 12,black_color, )
        self.add_text(prs, 0, Inches(2.55), Inches(5.64), Inches(1), Inches(1), seria, 12, black_color)

        # Slide 1

        module_1_score = f"{self.module_1}%"
        module_2_score = f"{self.module_2}%"
        module_3_score = f"{self.module_3}%"
        module_4_score = f"{self.module_4}%"
        module_5_score = f"{self.module_5}%"

        self.add_text(prs, 1, Inches(0.54), Inches(0.028), Inches(1), Inches(0.8), seria, 11, black_color)
        self.add_text(prs, 1, Inches(1), Inches(0.75), Inches(8), Inches(1), text, 28, text_color,
                      alignment=PP_ALIGN.CENTER)

        self.add_text(prs, 1, Inches(2.44), Inches(3.732283), Inches(1.7440944882), Inches(0.3346456693),
                      module_1_score, 14, black_color)  # Module 1 Score
        self.add_text(prs, 1, Inches(2.44), Inches(4.61), Inches(1.7440944882), Inches(0.3346456693),
                      module_2_score, 14, black_color)  # Module 2 Score
        self.add_text(prs, 1, Inches(2.44), Inches(5.55), Inches(1.7440944882), Inches(0.3346456693),
                      module_3_score, 14, black_color)  # Module 3 Score
        self.add_text(prs, 1, Inches(6.3385826772), Inches(3.732283), Inches(1.7440944882), Inches(0.3346456693),
                      module_4_score, 14, black_color)  # Module 4 Score
        self.add_text(prs, 1, Inches(6.3385826772), Inches(4.6102362), Inches(1.7440944882), Inches(0.3346456693),
                      module_5_score, 14, black_color)  # Module 5 Score
        self.add_image(prs, 1, qr_code_img_back, Inches(8.4448818898), Inches(4.8385826772), Inches(1.3))

        # add_image(prs, slide, png, left,top,height)

        # add_text(prs, slide, left, top, width, height, text, size, color)

        pptx_buffer = BytesIO()
        prs.save(pptx_buffer)
        pptx_buffer.seek(0)

        return pptx_buffer

    def overlay_qr_code_front(self, background_path, qr_code_path, output_path, position, qr_size):
        background = Image.open(background_path)
        qr_code = Image.open(qr_code_path).resize((qr_size, qr_size), Image.LANCZOS)

        qr_code = qr_code.convert("RGBA")
        data = qr_code.getdata()
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        qr_code.putdata(new_data)

        background.paste(qr_code, position, qr_code)

        font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=100)
        draw = ImageDraw.Draw(background)
        seria_font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=45)

        black_color = (0, 0, 0)
        text_color = (0x54, 0x30, 0xCE)

        text = f"{self.last_name} {self.first_name}  {self.middle_name}"
        draw.text((1500, 980), text, fill=text_color, anchor="ms", font=font)
        draw.text((800, 1810), f"{self.seria} {self.sertificate_id}", fill=black_color, font=seria_font)
        draw.text((1420, 1810), f"{self.issue_date}", fill=black_color, font=seria_font)
        background.save(output_path)

    def overlay_qr_code_back(self, background_path, qr_code_path, output_path, position, qr_size):
        background = Image.open(background_path)
        qr_code = Image.open(qr_code_path).resize((qr_size, qr_size), Image.LANCZOS)

        qr_code = qr_code.convert("RGBA")
        data = qr_code.getdata()
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        qr_code.putdata(new_data)

        background.paste(qr_code, position, qr_code)

        font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=100)
        draw = ImageDraw.Draw(background)
        seria_font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=50)
        seria_font_sr = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=45)

        black_color = (0, 0, 0)
        text_color = (0x54, 0x30, 0xCE)
        text = f"{self.last_name} {self.first_name}  {self.middle_name}"
        draw.text((1500, 440), text, fill=text_color, anchor="ms", font=font)
        draw.text((190, 128), f"{self.seria} {self.sertificate_id}", fill=black_color, font=seria_font_sr)
        draw.text((880, 1250), f"{self.module_1}%", fill=black_color, font=seria_font)
        draw.text((880, 1510), f"{self.module_2}%", fill=black_color, font=seria_font)
        draw.text((880, 1770), f"{self.module_3}%", fill=black_color, font=seria_font)
        draw.text((2100, 1250), f"{self.module_4}%", fill=black_color, font=seria_font)
        draw.text((2100, 1510), f"{self.module_5}%", fill=black_color, font=seria_font)

        background.save(output_path)

    def save(self, *args, **kwargs):
        if not self.sertificate_id:
            last_id = get_last_sertificate_id()
            new_id = last_id + 1
            self.sertificate_id = str(new_id).zfill(7)

        if not self.sertificate_id_numeric:
            last_numeric_id = get_last_sertificate_id_numeric()
            new_numeric_id = last_numeric_id + 1
            self.sertificate_id_numeric = new_numeric_id

        # Here, on which link you need to create a qr code, put that link
        qr_code_img = qrcode.make(f"https://certificate.astrum.uz/student/PC{self.sertificate_id}")
        canvas = Image.new('RGB', (380, 380), 'white')
        canvas.paste(qr_code_img)
        buffer = BytesIO()
        canvas.save(buffer, format='PNG')
        qr_code_file_name = f'qr_code-{self.seria}-{self.sertificate_id}.png'
        self.qr_code.save(qr_code_file_name, File(buffer), save=False)
        canvas.close()

        if not self.sertificate_front:
            background_image_path_front = os.path.join(settings.MEDIA_ROOT, 'template/php-front.png')
            qr_code_image_path_front = os.path.join(settings.MEDIA_ROOT,
                                                    f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
            output_image_path_front = os.path.join(settings.MEDIA_ROOT,
                                                   f'module_front/certificate-{self.seria}-{self.sertificate_id}.png')
            self.overlay_qr_code_front(background_image_path_front, qr_code_image_path_front, output_image_path_front,
                                       (115, 1435), 390)
            self.sertificate_front = os.path.relpath(output_image_path_front, settings.MEDIA_ROOT)

        if not self.sertificate_back:
            background_image_path_back = os.path.join(settings.MEDIA_ROOT, 'template/php-back.png')
            qr_code_image_path_back = os.path.join(settings.MEDIA_ROOT,
                                                   f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
            output_image_path_back = os.path.join(settings.MEDIA_ROOT,
                                                  f'module_back/certificate-{self.seria}-{self.sertificate_id}.png')
            self.overlay_qr_code_back(background_image_path_back, qr_code_image_path_back, output_image_path_back,
                                      (2535, 1445), 390)

            self.sertificate_back = output_image_path_back

        if not self.pptx_file:
            pptx_buffer = self.generate_certificate()
            self.pptx_file.save(
                f'{self.course}-{self.seria}-{self.sertificate_id}-{self.last_name}-{self.first_name}-{self.middle_name}.pptx',
                ContentFile(pptx_buffer.read()), save=False)

        super(PHPModule, self).save(*args, **kwargs)


class PythonModule(models.Model):
    first_name = models.CharField(max_length=50)
    last_name = models.CharField(max_length=50)
    middle_name = models.CharField(max_length=50)
    issue_date = models.CharField(max_length=30, blank=True)
    created_at = models.DateTimeField(auto_now_add=True)
    module_1 = models.CharField(max_length=5)
    module_2 = models.CharField(max_length=5)
    module_3 = models.CharField(max_length=5)
    module_4 = models.CharField(max_length=5)
    module_5 = models.CharField(max_length=5)
    certificate_link = models.CharField(max_length=250, blank=True)
    seria = models.CharField(max_length=3, default="PC")
    course = models.CharField(max_length=8, default="PYTHON")
    sertificate_id = models.CharField(max_length=7, unique=True, null=True, blank=True)
    sertificate_id_numeric = models.IntegerField(unique=True, null=True, blank=True)
    qr_code = models.ImageField(upload_to="module_qr_code", null=True, blank=True)
    pptx_file = models.FileField(upload_to="pptx_module", blank=True)
    sertificate_front = models.FileField(upload_to="module_front", blank=True)
    sertificate_back = models.FileField(upload_to="module_back", blank=True)


    def __str__(self):
        return self.first_name

    @staticmethod
    def add_image(prs, slide_index, image_path, left, top, height):
        slide = prs.slides[slide_index]
        image = Image.open(image_path)
        image = image.convert('RGBA')
        data = list(image.getdata())
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        new_image = Image.new('RGBA', image.size)
        new_image.putdata(new_data)
        image_stream = BytesIO()
        new_image.save(image_stream, format="PNG")
        image_stream.seek(0)
        slide.shapes.add_picture(image_stream, left, top, height=height)

    @staticmethod
    def add_text(prs, slide_index, left, top, width, height, input_text, font_size, font_color,
                 font_name='Gilroy', alignment=PP_ALIGN.LEFT):
        slide = prs.slides[slide_index]
        text_box = slide.shapes.add_textbox(left, top, width, height)
        text_frame = text_box.text_frame
        text = text_frame.add_paragraph()
        text.text = input_text
        text.font.size = Pt(font_size)
        text.font.name = font_name
        text.font.color.rgb = RGBColor(font_color[0], font_color[1], font_color[2])
        text.alignment = alignment
        for run in text.runs:
            run.font.bold = True

    def generate_certificate(self):
        pptx_template_path = os.path.join(settings.MEDIA_ROOT, 'template/python_module.pptx')
        prs = Presentation(pptx_template_path)

        # Slide 0

        black_color = (0, 0, 0,)
        text_color = (0x54, 0x30, 0xCE)
        text = f"{self.last_name} {self.first_name} {self.middle_name}"
        seria = f"{self.seria} {self.sertificate_id}"
        qr_code = os.path.join(settings.MEDIA_ROOT, f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
        qr_code_img_back = os.path.join(settings.MEDIA_ROOT, f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
        issued = f"{self.issue_date}"

        self.add_image(prs, 0, qr_code, Inches(0.3976377953), Inches(4.7952755906), Inches(1.3))
        self.add_text(prs, 0, Inches(1), Inches(2.55), Inches(8), Inches(1), text, 28, text_color,
                      alignment=PP_ALIGN.CENTER)
        self.add_text(prs, 0, Inches(4.645669), Inches(5.64), Inches(1), Inches(1), issued, 12,black_color, )
        self.add_text(prs, 0, Inches(2.55), Inches(5.64), Inches(1), Inches(1), seria, 12, black_color)

        # Slide 1

        module_1_score = f"{self.module_1} %"
        module_2_score = f"{self.module_2} %"
        module_3_score = f"{self.module_3} %"
        module_4_score = f"{self.module_4} %"
        module_5_score = f"{self.module_5} %"

        self.add_text(prs, 1, Inches(0.54), Inches(0.028), Inches(1), Inches(0.8), seria, 11, black_color)
        self.add_text(prs, 1, Inches(1), Inches(0.75), Inches(8), Inches(1), text, 28, text_color,
                      alignment=PP_ALIGN.CENTER)

        self.add_text(prs, 1, Inches(2.44), Inches(3.732283), Inches(1.7440944882), Inches(0.3346456693),
                      module_1_score, 14, black_color)  # Module 1 Score
        self.add_text(prs, 1, Inches(2.44), Inches(4.61), Inches(1.7440944882), Inches(0.3346456693),
                      module_2_score, 14, black_color)  # Module 2 Score
        self.add_text(prs, 1, Inches(2.44), Inches(5.55), Inches(1.7440944882), Inches(0.3346456693),
                      module_3_score, 14, black_color)  # Module 3 Score
        self.add_text(prs, 1, Inches(6.3385826772), Inches(3.732283), Inches(1.7440944882), Inches(0.3346456693),
                      module_4_score, 14, black_color)  # Module 4 Score
        self.add_text(prs, 1, Inches(6.3385826772), Inches(4.6102362), Inches(1.7440944882), Inches(0.3346456693),
                      module_5_score, 14, black_color)  # Module 5 Score
        self.add_image(prs, 1, qr_code_img_back, Inches(8.4448818898), Inches(4.8385826772), Inches(1.3))

        # add_image(prs, slide, png, left,top,height)

        # add_text(prs, slide, left, top, width, height, text, size, color)

        pptx_buffer = BytesIO()
        prs.save(pptx_buffer)
        pptx_buffer.seek(0)

        return pptx_buffer

    def overlay_qr_code_front(self, background_path, qr_code_path, output_path, position, qr_size):
        background = Image.open(background_path)
        qr_code = Image.open(qr_code_path).resize((qr_size, qr_size), Image.LANCZOS)

        qr_code = qr_code.convert("RGBA")
        data = qr_code.getdata()
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        qr_code.putdata(new_data)

        background.paste(qr_code, position, qr_code)

        font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=100)
        draw = ImageDraw.Draw(background)
        seria_font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=45)

        black_color = (0, 0, 0)
        text_color = (0x54, 0x30, 0xCE)

        text = f"{self.last_name} {self.first_name}  {self.middle_name}"
        draw.text((1500, 980), text, fill=text_color, anchor="ms", font=font)
        draw.text((800, 1810), f"{self.seria} {self.sertificate_id}", fill=black_color, font=seria_font)
        draw.text((1420, 1810), f"{self.issue_date}", fill=black_color, font=seria_font)
        background.save(output_path)

    def overlay_qr_code_back(self, background_path, qr_code_path, output_path, position, qr_size):
        background = Image.open(background_path)
        qr_code = Image.open(qr_code_path).resize((qr_size, qr_size), Image.LANCZOS)

        qr_code = qr_code.convert("RGBA")
        data = qr_code.getdata()
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        qr_code.putdata(new_data)

        background.paste(qr_code, position, qr_code)

        font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=100)
        draw = ImageDraw.Draw(background)
        seria_font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=50)
        seria_font_sr = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=45)

        black_color = (0, 0, 0)
        text_color = (0x54, 0x30, 0xCE)
        text = f"{self.last_name} {self.first_name}  {self.middle_name}"
        draw.text((1500, 440), text, fill=text_color, anchor="ms", font=font)
        draw.text((190, 128), f"{self.seria} {self.sertificate_id}", fill=black_color, font=seria_font_sr)
        draw.text((880, 1250), f"{self.module_1} %", fill=black_color, font=seria_font)
        draw.text((880, 1510), f"{self.module_2} %", fill=black_color, font=seria_font)
        draw.text((880, 1770), f"{self.module_3} %", fill=black_color, font=seria_font)
        draw.text((2100, 1250), f"{self.module_4} %", fill=black_color, font=seria_font)
        draw.text((2100, 1510), f"{self.module_5} %", fill=black_color, font=seria_font)

        background.save(output_path)

    def save(self, *args, **kwargs):

        if not self.sertificate_id:
            last_id = get_last_sertificate_id()
            new_id = last_id + 1
            self.sertificate_id = str(new_id).zfill(7)

        if not self.sertificate_id_numeric:
            last_numeric_id = get_last_sertificate_id_numeric()
            new_numeric_id = last_numeric_id + 1
            self.sertificate_id_numeric = new_numeric_id

        # Here, on which link you need to create a qr code, put that link
        qr_code_img = qrcode.make(f"https://certificate.astrum.uz/student/PC{self.sertificate_id}")
        canvas = Image.new('RGB', (380, 380), 'white')
        canvas.paste(qr_code_img)
        buffer = BytesIO()
        canvas.save(buffer, format='PNG')
        qr_code_file_name = f'qr_code-{self.seria}-{self.sertificate_id}.png'
        self.qr_code.save(qr_code_file_name, File(buffer), save=False)
        canvas.close()

        if not self.sertificate_front:
            background_image_path_front = os.path.join(settings.MEDIA_ROOT, 'template/python-front.png')
            qr_code_image_path_front = os.path.join(settings.MEDIA_ROOT,
                                                    f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
            output_image_path_front = os.path.join(settings.MEDIA_ROOT,
                                                   f'module_front/certificate-{self.seria}-{self.sertificate_id}.png')
            self.overlay_qr_code_front(background_image_path_front, qr_code_image_path_front, output_image_path_front,
                                       (115, 1435), 390)
            self.sertificate_front = os.path.relpath(output_image_path_front, settings.MEDIA_ROOT)

        if not self.sertificate_back:
            background_image_path_back = os.path.join(settings.MEDIA_ROOT, 'template/python-back.png')
            qr_code_image_path_back = os.path.join(settings.MEDIA_ROOT,
                                                   f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
            output_image_path_back = os.path.join(settings.MEDIA_ROOT,
                                                  f'module_back/certificate-{self.seria}-{self.sertificate_id}.png')
            self.overlay_qr_code_back(background_image_path_back, qr_code_image_path_back, output_image_path_back,
                                      (2535, 1445), 390)

            self.sertificate_back = output_image_path_back

        if not self.pptx_file:
            pptx_buffer = self.generate_certificate()
            self.pptx_file.save(
                f'{self.course}-{self.seria}-{self.sertificate_id}-{self.last_name}-{self.first_name}-{self.middle_name}.pptx',
                ContentFile(pptx_buffer.read()), save=False)

        super(PythonModule, self).save(*args, **kwargs)

class ReactModule(models.Model):
    first_name = models.CharField(max_length=50)
    last_name = models.CharField(max_length=50)
    middle_name = models.CharField(max_length=50)
    issue_date = models.CharField(max_length=30, blank=True)
    created_at = models.DateTimeField(auto_now_add=True)
    module_1 = models.CharField(max_length=5)
    module_2 = models.CharField(max_length=5)
    module_3 = models.CharField(max_length=5)
    module_4 = models.CharField(max_length=5)
    module_5 = models.CharField(max_length=5)
    certificate_link = models.CharField(max_length=250, blank=True)
    seria = models.CharField(max_length=3, default="PC")
    course = models.CharField(max_length=8  , default="REACT")
    sertificate_id = models.CharField(max_length=7, unique=True, null=True, blank=True)
    sertificate_id_numeric = models.IntegerField(unique=True, null=True, blank=True)
    qr_code = models.ImageField(upload_to="module_qr_code", null=True, blank=True)
    pptx_file = models.FileField(upload_to="pptx_module", blank=True)
    sertificate_front = models.FileField(upload_to="module_front", blank=True)
    sertificate_back = models.FileField(upload_to="module_back", blank=True)

    def __str__(self):
        return self.first_name

    @staticmethod
    def add_image(prs, slide_index, image_path, left, top, height):
        slide = prs.slides[slide_index]
        image = Image.open(image_path)
        image = image.convert('RGBA')
        data = list(image.getdata())
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        new_image = Image.new('RGBA', image.size)
        new_image.putdata(new_data)
        image_stream = BytesIO()
        new_image.save(image_stream, format="PNG")
        image_stream.seek(0)
        slide.shapes.add_picture(image_stream, left, top, height=height)

    @staticmethod
    def add_text(prs, slide_index, left, top, width, height, input_text, font_size, font_color,
                 font_name='Gilroy', alignment=PP_ALIGN.LEFT):
        slide = prs.slides[slide_index]
        text_box = slide.shapes.add_textbox(left, top, width, height)
        text_frame = text_box.text_frame
        text = text_frame.add_paragraph()
        text.text = input_text
        text.font.size = Pt(font_size)
        text.font.name = font_name
        text.font.color.rgb = RGBColor(font_color[0], font_color[1], font_color[2])
        text.alignment = alignment
        for run in text.runs:
            run.font.bold = True

    def generate_certificate(self):
        pptx_template_path = os.path.join(settings.MEDIA_ROOT, 'template/react.pptx')
        prs = Presentation(pptx_template_path)

        # Slide 0

        black_color = (0, 0, 0,)
        text_color = (0x54, 0x30, 0xCE)
        text = f"{self.last_name} {self.first_name} {self.middle_name}"
        seria = f"{self.seria} {self.sertificate_id}"
        qr_code = os.path.join(settings.MEDIA_ROOT, f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
        qr_code_img_back = os.path.join(settings.MEDIA_ROOT,
                                        f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
        issued = f"{self.issue_date}"

        self.add_image(prs, 0, qr_code, Inches(0.3976377953), Inches(4.7952755906), Inches(1.3))
        self.add_text(prs, 0, Inches(1), Inches(2.55), Inches(8), Inches(1), text, 28, text_color,
                      alignment=PP_ALIGN.CENTER)
        self.add_text(prs, 0, Inches(4.645669), Inches(5.64), Inches(1), Inches(1), issued, 12,black_color, )
        self.add_text(prs, 0, Inches(2.55), Inches(5.64), Inches(1), Inches(1), seria, 12, black_color)

        # Slide 1

        module_1_score = f"{self.module_1} %"
        module_2_score = f"{self.module_2} %"
        module_3_score = f"{self.module_3} %"
        module_4_score = f"{self.module_4} %"
        module_5_score = f"{self.module_5} %"

        self.add_text(prs, 1, Inches(0.54), Inches(0.028), Inches(1), Inches(0.8), seria, 11, black_color)
        self.add_text(prs, 1, Inches(1), Inches(0.75), Inches(8), Inches(1), text, 28, text_color,
                      alignment=PP_ALIGN.CENTER)

        self.add_text(prs, 1, Inches(2.44), Inches(3.732283), Inches(1.7440944882), Inches(0.3346456693),
                      module_1_score, 14, black_color)  # Module 1 Score
        self.add_text(prs, 1, Inches(2.44), Inches(4.61), Inches(1.7440944882), Inches(0.3346456693),
                      module_2_score, 14, black_color)  # Module 2 Score
        self.add_text(prs, 1, Inches(2.44), Inches(5.55), Inches(1.7440944882), Inches(0.3346456693),
                      module_3_score, 14, black_color)  # Module 3 Score
        self.add_text(prs, 1, Inches(6.3385826772), Inches(3.732283), Inches(1.7440944882), Inches(0.3346456693),
                      module_4_score, 14, black_color)  # Module 4 Score
        self.add_text(prs, 1, Inches(6.3385826772), Inches(4.6102362), Inches(1.7440944882), Inches(0.3346456693),
                      module_5_score, 14, black_color)  # Module 5 Score
        self.add_image(prs, 1, qr_code_img_back, Inches(8.4448818898), Inches(4.8385826772), Inches(1.3))

        # add_image(prs, slide, png, left,top,height)

        # add_text(prs, slide, left, top, width, height, text, size, color)

        pptx_buffer = BytesIO()
        prs.save(pptx_buffer)
        pptx_buffer.seek(0)

        return pptx_buffer

    def overlay_qr_code_front(self, background_path, qr_code_path, output_path, position, qr_size):
        background = Image.open(background_path)
        qr_code = Image.open(qr_code_path).resize((qr_size, qr_size), Image.LANCZOS)

        qr_code = qr_code.convert("RGBA")
        data = qr_code.getdata()
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        qr_code.putdata(new_data)

        background.paste(qr_code, position, qr_code)

        font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=100)
        draw = ImageDraw.Draw(background)
        seria_font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=45)

        black_color = (0, 0, 0)
        text_color = (0x54, 0x30, 0xCE)

        text = f"{self.last_name} {self.first_name}  {self.middle_name}"
        draw.text((1500, 980), text, fill=text_color, anchor="ms", font=font)
        draw.text((800, 1810), f"{self.seria} {self.sertificate_id}", fill=black_color, font=seria_font)
        draw.text((1420, 1810), f"{self.issue_date}", fill=black_color, font=seria_font)
        background.save(output_path)

    def overlay_qr_code_back(self, background_path, qr_code_path, output_path, position, qr_size):
        background = Image.open(background_path)
        qr_code = Image.open(qr_code_path).resize((qr_size, qr_size), Image.LANCZOS)

        qr_code = qr_code.convert("RGBA")
        data = qr_code.getdata()
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        qr_code.putdata(new_data)

        background.paste(qr_code, position, qr_code)

        font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=100)
        draw = ImageDraw.Draw(background)
        seria_font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=50)
        seria_font_sr = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=45)

        black_color = (0, 0, 0)
        text_color = (0x54, 0x30, 0xCE)
        text = f"{self.last_name} {self.first_name}  {self.middle_name}"
        draw.text((1500, 440), text, fill=text_color, anchor="ms", font=font)
        draw.text((190, 128), f"{self.seria} {self.sertificate_id}", fill=black_color, font=seria_font_sr)
        draw.text((880, 1250), f"{self.module_1} %", fill=black_color, font=seria_font)
        draw.text((880, 1510), f"{self.module_2} %", fill=black_color, font=seria_font)
        draw.text((880, 1770), f"{self.module_3} %", fill=black_color, font=seria_font)
        draw.text((2100, 1250), f"{self.module_4} %", fill=black_color, font=seria_font)
        draw.text((2100, 1510), f"{self.module_5} %", fill=black_color, font=seria_font)

        background.save(output_path)

    def save(self, *args, **kwargs):
        if not self.sertificate_id:
            last_id = get_last_sertificate_id()
            new_id = last_id + 1
            self.sertificate_id = str(new_id).zfill(7)

        if not self.sertificate_id_numeric:
            last_numeric_id = get_last_sertificate_id_numeric()
            new_numeric_id = last_numeric_id + 1
            self.sertificate_id_numeric = new_numeric_id

        # Here, on which link you need to create a qr code, put that link
        qr_code_img = qrcode.make(f"https://certificate.astrum.uz/student/PC{self.sertificate_id}")
        canvas = Image.new('RGB', (380, 380), 'white')
        canvas.paste(qr_code_img)
        buffer = BytesIO()
        canvas.save(buffer, format='PNG')
        qr_code_file_name = f'qr_code-{self.seria}-{self.sertificate_id}.png'
        self.qr_code.save(qr_code_file_name, File(buffer), save=False)
        canvas.close()

        if not self.sertificate_front:
            background_image_path_front = os.path.join(settings.MEDIA_ROOT, 'template/react-front.png')
            qr_code_image_path_front = os.path.join(settings.MEDIA_ROOT,
                                                    f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
            output_image_path_front = os.path.join(settings.MEDIA_ROOT,
                                                   f'module_front/certificate-{self.seria}-{self.sertificate_id}.png')
            self.overlay_qr_code_front(background_image_path_front, qr_code_image_path_front, output_image_path_front,
                                       (115, 1435), 390)
            self.sertificate_front = os.path.relpath(output_image_path_front, settings.MEDIA_ROOT)

        if not self.sertificate_back:
            background_image_path_back = os.path.join(settings.MEDIA_ROOT, 'template/react-back.png')
            qr_code_image_path_back = os.path.join(settings.MEDIA_ROOT,
                                                   f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
            output_image_path_back = os.path.join(settings.MEDIA_ROOT,
                                                  f'module_back/certificate-{self.seria}-{self.sertificate_id}.png')
            self.overlay_qr_code_back(background_image_path_back, qr_code_image_path_back, output_image_path_back,
                                      (2535, 1445), 390)

            self.sertificate_back = output_image_path_back

        if not self.pptx_file:
            pptx_buffer = self.generate_certificate()
            self.pptx_file.save(
                f'{self.course}-{self.seria}-{self.sertificate_id}-{self.last_name}-{self.first_name}-{self.middle_name}.pptx',
                ContentFile(pptx_buffer.read()), save=False)

        super(ReactModule, self).save(*args, **kwargs)


class CSharpModule(models.Model):
    first_name = models.CharField(max_length=50)
    last_name = models.CharField(max_length=50)
    middle_name = models.CharField(max_length=50)
    issue_date = models.CharField(max_length=30, blank=True)
    created_at = models.DateTimeField(auto_now_add=True)
    module_1 = models.CharField(max_length=5)
    module_2 = models.CharField(max_length=5)
    module_3 = models.CharField(max_length=5)
    module_4 = models.CharField(max_length=5)
    module_5 = models.CharField(max_length=5)
    certificate_link = models.CharField(max_length=250, blank=True)
    seria = models.CharField(max_length=3, default="PC")
    course = models.CharField(max_length=3, default="C#")
    sertificate_id = models.CharField(max_length=7, unique=True, null=True, blank=True)
    sertificate_id_numeric = models.IntegerField(unique=True, null=True, blank=True)
    qr_code = models.ImageField(upload_to="module_qr_code", null=True, blank=True)
    pptx_file = models.FileField(upload_to="pptx_module", blank=True)
    sertificate_front = models.FileField(upload_to="module_front", blank=True)
    sertificate_back = models.FileField(upload_to="module_back", blank=True)

    def __str__(self):
        return self.first_name


    @staticmethod
    def add_image(prs, slide_index, image_path, left, top, height):
        slide = prs.slides[slide_index]
        image = Image.open(image_path)
        image = image.convert('RGBA')
        data = list(image.getdata())
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        new_image = Image.new('RGBA', image.size)
        new_image.putdata(new_data)
        image_stream = BytesIO()
        new_image.save(image_stream, format="PNG")
        image_stream.seek(0)
        slide.shapes.add_picture(image_stream, left, top, height=height)

    @staticmethod
    def add_text(prs, slide_index, left, top, width, height, input_text, font_size, font_color,
                 font_name='Gilroy', alignment=PP_ALIGN.LEFT):
        slide = prs.slides[slide_index]
        text_box = slide.shapes.add_textbox(left, top, width, height)
        text_frame = text_box.text_frame
        text = text_frame.add_paragraph()
        text.text = input_text
        text.font.size = Pt(font_size)
        text.font.name = font_name
        text.font.color.rgb = RGBColor(font_color[0], font_color[1], font_color[2])
        text.alignment = alignment
        for run in text.runs:
            run.font.bold = True

    def generate_certificate(self):
        pptx_template_path = os.path.join(settings.MEDIA_ROOT, 'template/csh.pptx')
        prs = Presentation(pptx_template_path)

        # Slide 0

        black_color = (0, 0, 0,)
        text_color = (0x54, 0x30, 0xCE)
        text = f"{self.last_name} {self.first_name} {self.middle_name}"
        seria = f"{self.seria} {self.sertificate_id}"
        qr_code = os.path.join(settings.MEDIA_ROOT, f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
        qr_code_img_back = os.path.join(settings.MEDIA_ROOT,
                                        f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
        issued = f"{self.issue_date}"

        self.add_image(prs, 0, qr_code, Inches(0.3976377953), Inches(4.7952755906), Inches(1.3))
        self.add_text(prs, 0, Inches(1), Inches(2.55), Inches(8), Inches(1), text, 28, text_color,
                      alignment=PP_ALIGN.CENTER)
        self.add_text(prs, 0, Inches(4.645669), Inches(5.64), Inches(1), Inches(1), issued, 12,black_color, )
        self.add_text(prs, 0, Inches(2.55), Inches(5.64), Inches(1), Inches(1), seria, 12, black_color)

        # Slide 1

        module_1_score = f"{self.module_1} %"
        module_2_score = f"{self.module_2} %"
        module_3_score = f"{self.module_3} %"
        module_4_score = f"{self.module_4} %"
        module_5_score = f"{self.module_5} %"

        self.add_text(prs, 1, Inches(0.54), Inches(0.028), Inches(1), Inches(0.8), seria, 11, black_color)
        self.add_text(prs, 1, Inches(1), Inches(0.75), Inches(8), Inches(1), text, 28, text_color,
                      alignment=PP_ALIGN.CENTER)

        self.add_text(prs, 1, Inches(2.44), Inches(3.732283), Inches(1.7440944882), Inches(0.3346456693),
                      module_1_score, 14, black_color)  # Module 1 Score
        self.add_text(prs, 1, Inches(2.44), Inches(4.61), Inches(1.7440944882), Inches(0.3346456693),
                      module_2_score, 14, black_color)  # Module 2 Score
        self.add_text(prs, 1, Inches(2.44), Inches(5.55), Inches(1.7440944882), Inches(0.3346456693),
                      module_3_score, 14, black_color)  # Module 3 Score
        self.add_text(prs, 1, Inches(6.3385826772), Inches(3.732283), Inches(1.7440944882), Inches(0.3346456693),
                      module_4_score, 14, black_color)  # Module 4 Score
        self.add_text(prs, 1, Inches(6.3385826772), Inches(4.6102362), Inches(1.7440944882), Inches(0.3346456693),
                      module_5_score, 14, black_color)  # Module 5 Score
        self.add_image(prs, 1, qr_code_img_back, Inches(8.4448818898), Inches(4.8385826772), Inches(1.3))

        # add_image(prs, slide, png, left,top,height)

        # add_text(prs, slide, left, top, width, height, text, size, color)

        pptx_buffer = BytesIO()
        prs.save(pptx_buffer)
        pptx_buffer.seek(0)

        return pptx_buffer

    def overlay_qr_code_front(self, background_path, qr_code_path, output_path, position, qr_size):
        background = Image.open(background_path)
        qr_code = Image.open(qr_code_path).resize((qr_size, qr_size), Image.LANCZOS)

        qr_code = qr_code.convert("RGBA")
        data = qr_code.getdata()
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        qr_code.putdata(new_data)

        background.paste(qr_code, position, qr_code)

        font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=100)
        draw = ImageDraw.Draw(background)
        seria_font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=45)

        black_color = (0, 0, 0)
        text_color = (0x54, 0x30, 0xCE)

        text = f"{self.last_name} {self.first_name}  {self.middle_name}"
        draw.text((1500, 980), text, fill=text_color, anchor="ms", font=font)
        draw.text((800, 1810), f"{self.seria} {self.sertificate_id}", fill=black_color, font=seria_font)
        draw.text((1420, 1810), f"{self.issue_date}", fill=black_color, font=seria_font)
        background.save(output_path)

    def overlay_qr_code_back(self, background_path, qr_code_path, output_path, position, qr_size):
        background = Image.open(background_path)
        qr_code = Image.open(qr_code_path).resize((qr_size, qr_size), Image.LANCZOS)

        qr_code = qr_code.convert("RGBA")
        data = qr_code.getdata()
        new_data = [(255, 255, 255, 0) if item[:3] == (255, 255, 255) else item for item in data]
        qr_code.putdata(new_data)

        background.paste(qr_code, position, qr_code)

        font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=100)
        draw = ImageDraw.Draw(background)
        seria_font = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=50)
        seria_font_sr = ImageFont.truetype(os.path.join(settings.MEDIA_ROOT, 'template/Gilroy-Black.ttf'), size=45)

        black_color = (0, 0, 0)
        text_color = (0x54, 0x30, 0xCE)
        text = f"{self.last_name} {self.first_name}  {self.middle_name}"
        draw.text((1500, 440), text, fill=text_color, anchor="ms", font=font)
        draw.text((190, 128), f"{self.seria} {self.sertificate_id}", fill=black_color, font=seria_font_sr)
        draw.text((880, 1250), f"{self.module_1} %", fill=black_color, font=seria_font)
        draw.text((880, 1510), f"{self.module_2} %", fill=black_color, font=seria_font)
        draw.text((880, 1770), f"{self.module_3} %", fill=black_color, font=seria_font)
        draw.text((2100, 1250), f"{self.module_4} %", fill=black_color, font=seria_font)
        draw.text((2100, 1510), f"{self.module_5} %", fill=black_color, font=seria_font)

        background.save(output_path)

    def save(self, *args, **kwargs):
        if not self.sertificate_id:
            last_id = get_last_sertificate_id()
            new_id = last_id + 1
            self.sertificate_id = str(new_id).zfill(7)

        if not self.sertificate_id_numeric:
            last_numeric_id = get_last_sertificate_id_numeric()
            new_numeric_id = last_numeric_id + 1
            self.sertificate_id_numeric = new_numeric_id

        # Here, on which link you need to create a qr code, put that link
        qr_code_img = qrcode.make(f"https://certificate.astrum.uz/student/PC{self.sertificate_id}")
        canvas = Image.new('RGB', (380, 380), 'white')
        canvas.paste(qr_code_img)
        buffer = BytesIO()
        canvas.save(buffer, format='PNG')
        qr_code_file_name = f'qr_code-{self.seria}-{self.sertificate_id}.png'
        self.qr_code.save(qr_code_file_name, File(buffer), save=False)
        canvas.close()

        if not self.sertificate_front:
            background_image_path_front = os.path.join(settings.MEDIA_ROOT, 'template/csh-front.png')
            qr_code_image_path_front = os.path.join(settings.MEDIA_ROOT,
                                                    f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
            output_image_path_front = os.path.join(settings.MEDIA_ROOT,
                                                   f'module_front/certificate-{self.seria}-{self.sertificate_id}.png')
            self.overlay_qr_code_front(background_image_path_front, qr_code_image_path_front, output_image_path_front,
                                       (115, 1435), 390)
            self.sertificate_front = os.path.relpath(output_image_path_front, settings.MEDIA_ROOT)

        if not self.sertificate_back:
            background_image_path_back = os.path.join(settings.MEDIA_ROOT, 'template/csh-back.png')
            qr_code_image_path_back = os.path.join(settings.MEDIA_ROOT,
                                                   f'module_qr_code/qr_code-{self.seria}-{self.sertificate_id}.png')
            output_image_path_back = os.path.join(settings.MEDIA_ROOT,
                                                  f'module_back/certificate-{self.seria}-{self.sertificate_id}.png')
            self.overlay_qr_code_back(background_image_path_back, qr_code_image_path_back, output_image_path_back,
                                      (2535, 1445), 390)

            self.sertificate_back = output_image_path_back

        if not self.pptx_file:
            pptx_buffer = self.generate_certificate()
            self.pptx_file.save(
                f'{self.course}-{self.seria}-{self.sertificate_id}-{self.last_name}-{self.first_name}-{self.middle_name}.pptx',
                ContentFile(pptx_buffer.read()), save=False)

        super(CSharpModule, self).save(*args, **kwargs)
